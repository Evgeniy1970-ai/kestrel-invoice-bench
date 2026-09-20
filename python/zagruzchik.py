#!/usr/bin/env python3
"""
Kestrel Invoice Bench — загрузка справочников и эталона.

Что делает:
  kestrel-price-list.xlsx        -> схема ref   (заведения, поставщики, каталог, версии, цены)
  kestrel-april-2026-batch.xlsx  -> схема truth (документы партии, журнал дефектов, сводка)

Запускается один раз, под владельцем базы. Роль bench_pipeline здесь не годится:
она справочники только читает, а к схеме truth доступа не имеет вовсе — ровно так,
как задумано.

Использование:
    pip install openpyxl psycopg2-binary

    # проверка разбора таблиц без базы
    python zagruzchik.py --dry-run \
        --price ../../02-dannye/kestrel-price-list.xlsx \
        --batch ../../03-partiya/kestrel-april-2026-batch.xlsx

    # загрузка
    export DATABASE_URL="postgresql://postgres:ПАРОЛЬ@ХОСТ:5432/postgres"
    python zagruzchik.py \
        --price ../../02-dannye/kestrel-price-list.xlsx \
        --batch ../../03-partiya/kestrel-april-2026-batch.xlsx

Всё делается одной транзакцией: либо загрузилось целиком, либо не изменилось ничего.
Частично залитый справочник хуже пустого — он выглядит рабочим.
"""

import argparse
import os
import sys
from decimal import Decimal

from openpyxl import load_workbook

# ---------------------------------------------------------------------
# Того, чего нет в таблицах
#
# Порядок выставления счетов описан в potok-postavok.md, но в прайс-лист
# не попал.
#
# Макет здесь НЕ задаётся. Первая версия загрузчика брала его из
# maket-scheta.md, и по четырём поставщикам из двенадцати он разошёлся
# с партией — 168 документов из 260. Единственный источник правды о
# макете — сама партия: там лежат готовые PDF. Раскладка читается из
# truth.documents, см. layouts_from_batch().
# ---------------------------------------------------------------------
BILLING_MODE = {
    "VEG": "per_delivery", "MEA": "per_delivery", "DAI": "per_delivery",
    "BAK": "per_delivery", "BEV": "per_delivery",
    "DRY": "per_period", "LIN": "per_period", "COF": "per_period",
    "PAC": "per_period", "CHM": "per_period", "WST": "per_period",
    "EQP": "per_period",
}

INVOICE_NO_PATTERN = {
    "VEG": "GP-YYYY-MM-NNNN",
    "MEA": "HMF/YY/MM/NNNN",
    "DAI": "LD-NNNNNN",
    "DRY": "CDG-YYYY-NNNN",
    "BAK": "ABYY-MM-NNNN",
    "COF": "TCR-YYYY-NNNN",
    "BEV": "NB-YYYY-MM-NNNN",
    "PAC": "PP/YYYY/NNNN",
    "CHM": "VH-YY-NNNN",
    "LIN": "WLS-YYYY-MM-NNN",
    "WST": "RW-YYYY-NNNN",
    "EQP": "MES-YY-NNNN",
}

DUPLICATE_KINDS = {"same_file", "rebuilt_file", "new_number"}


# ---------------------------------------------------------------------
# Мелочи разбора
# ---------------------------------------------------------------------
def rows(wb, sheet):
    """Строки листа без заголовка, пустые хвосты отброшены."""
    data = list(wb[sheet].iter_rows(values_only=True))
    if not data:
        raise SystemExit(f"Лист {sheet} пуст")
    return [r for r in data[1:] if any(c is not None for c in r)]


def header(wb, sheet):
    return list(next(wb[sheet].iter_rows(values_only=True)))


def check_header(wb, sheet, expected):
    """Столбцы должны быть теми же. Иначе загрузка тихо перепутает поля."""
    got = header(wb, sheet)[: len(expected)]
    if got != expected:
        raise SystemExit(
            f"Лист {sheet}: столбцы не те, что ожидались.\n"
            f"  ожидалось: {expected}\n"
            f"  в файле:   {got}\n"
            f"Похоже, это другая версия файла. Сверьтесь с 00-opis.md."
        )


def dec(v):
    """Числа только через строку: float напрямую даёт 0.10000000000000001."""
    return None if v is None else Decimal(str(v))


def day(v):
    return None if v is None else (v.date() if hasattr(v, "date") else v)


def text(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


# ---------------------------------------------------------------------
# Прайс-лист -> ref
# ---------------------------------------------------------------------
def read_price_list(path, layouts):
    wb = load_workbook(path, read_only=True, data_only=True)

    check_header(wb, "Venues", ["Venue code", "Venue name", "Type", "Address"])
    check_header(wb, "Suppliers", ["Supplier code", "Supplier name", "Category",
                                   "Delivery frequency", "Payment terms", "Billing e-mail"])
    check_header(wb, "Catalog", ["SKU", "Supplier", "Description", "Unit", "Alt unit",
                                 "Alt unit factor", "VAT rate",
                                 "Discount threshold (qty)", "Discount %"])
    check_header(wb, "Prices", ["SKU", "Description", "Unit", "Unit price (EUR)",
                                "Valid from", "Valid to", "Price list version"])
    check_header(wb, "Price versions", ["Version", "Effective from", "Effective to", "Note"])

    venues = [
        (text(r[0]), text(r[1]), text(r[2]), text(r[3]))
        for r in rows(wb, "Venues")
    ]

    suppliers = []
    for r in rows(wb, "Suppliers"):
        code = text(r[0])
        name = text(r[1])
        if code not in BILLING_MODE:
            raise SystemExit(f"Поставщик {code} не описан в BILLING_MODE — поправьте загрузчик")
        if name not in layouts:
            raise SystemExit(
                f"Поставщика «{name}» нет в партии — макет взять неоткуда.\n"
                f"Скорее всего название в прайсе и в партии записано по-разному."
            )
        suppliers.append((
            code, name, text(r[2]), text(r[5]), text(r[4]),
            layouts[name], BILLING_MODE[code], INVOICE_NO_PATTERN[code], "approved",
        ))

    catalog = [
        (text(r[0]), text(r[1]), text(r[2]), text(r[3]), text(r[4]),
         dec(r[5]), dec(r[6]), dec(r[7]), dec(r[8]))
        for r in rows(wb, "Catalog")
    ]
    sku_supplier = {c[0]: c[1] for c in catalog}

    prices = [
        (text(r[0]), text(r[6]), text(r[2]), dec(r[3]), day(r[4]), day(r[5]))
        for r in rows(wb, "Prices")
    ]

    # Область действия версии выводится из данных, а не задаётся руками:
    # если все позиции версии принадлежат одному поставщику — это его версия.
    by_version = {}
    for sku, version, *_ in prices:
        by_version.setdefault(version, set()).add(sku_supplier[sku])

    versions = []
    for r in rows(wb, "Price versions"):
        code = text(r[0])
        owners = by_version.get(code, set())
        scope = owners.pop() if len(owners) == 1 else None
        versions.append((code, day(r[1]), day(r[2]), scope, text(r[3])))

    wb.close()
    return venues, suppliers, catalog, versions, prices


# ---------------------------------------------------------------------
# Партия -> truth
# ---------------------------------------------------------------------
def read_batch(path):
    wb = load_workbook(path, read_only=True, data_only=True)

    check_header(wb, "Documents", ["ID", "Номер счёта", "Поставщик", "Заведение", "Макет",
                                   "Получен", "Строк", "Итого EUR", "Вид файла", "Читаемость",
                                   "Смягчён", "Дефекты", "Дубль документа", "Вид дубля"])
    check_header(wb, "Journal", ["ID", "Номер счёта", "Дефект", "Поставщик", "Заведение",
                                 "Макет", "Скан", "Дата счёта", "Строка", "Поле",
                                 "Правильно", "В счёте", "Влияние EUR", "Пояснение"])

    documents = []
    for r in rows(wb, "Documents"):
        kind = text(r[13])
        if kind is not None and kind not in DUPLICATE_KINDS:
            raise SystemExit(f"{text(r[0])}: неизвестная разновидность дубля «{kind}»")
        documents.append((
            text(r[0]), text(r[1]), text(r[2]), text(r[3]), text(r[4]),
            day(r[5]), int(r[6]), dec(r[7]), text(r[8]), dec(r[9]),
            text(r[10]) == "да", text(r[11]), text(r[12]), kind,
        ))

    # Дубли ссылаются на оригиналы. Оригиналы вставляются первыми,
    # иначе внешний ключ на самого себя не даст вставить строку.
    documents.sort(key=lambda d: (d[13] is not None, d[0]))

    ids = {d[0] for d in documents}
    for d in documents:
        if d[12] and d[12] not in ids:
            raise SystemExit(f"{d[0]}: оригинал {d[12]} в партии не найден")

    defects = [
        (text(r[0]), text(r[1]), text(r[2]), text(r[3]), text(r[4]), text(r[5]),
         text(r[6]), day(r[7]), int(r[8]) if r[8] is not None else None,
         text(r[9]), text(r[10]), text(r[11]), dec(r[12]), text(r[13]))
        for r in rows(wb, "Journal")
    ]
    for d in defects:
        if d[0] not in ids:
            raise SystemExit(f"Журнал ссылается на документ {d[0]}, которого нет в партии")

    batch = [
        (text(r[0]), text(r[1]))
        for r in rows(wb, "Summary")
        if text(r[0]) is not None and text(r[1]) is not None
    ]

    # Раскладка макетов по поставщикам — из партии, а не из описания.
    layouts = {}
    for d in documents:
        supplier, layout = d[2], d[4]
        if layouts.setdefault(supplier, layout) != layout:
            raise SystemExit(f"У поставщика «{supplier}» в партии больше одного макета")

    wb.close()
    return documents, defects, batch, layouts


# ---------------------------------------------------------------------
# Запись
# ---------------------------------------------------------------------
def load(conn, price_data, batch_data):
    import psycopg2.extras as extras

    venues, suppliers, catalog, versions, prices = price_data
    documents, defects, batch, _ = batch_data

    with conn.cursor() as cur:
        cur.execute("select count(*) from truth.documents")
        if cur.fetchone()[0]:
            raise SystemExit(
                "В truth.documents уже есть строки. Эталон заливается один раз.\n"
                "Если это осознанная перезаливка, снимите заморозку:\n"
                "  alter table truth.documents disable trigger truth_documents_frozen;\n"
                "  alter table truth.defects   disable trigger truth_defects_frozen;\n"
                "  truncate truth.defects, truth.documents, truth.batch;\n"
                "и не забудьте включить триггеры обратно."
            )

        # Справочники перезаливаются целиком: они меняются, эталон — нет.
        cur.execute("delete from ref.prices")
        cur.execute("delete from ref.price_versions")
        cur.execute("delete from ref.sku_map")
        cur.execute("delete from ref.catalog_items")
        cur.execute("delete from ref.suppliers")
        cur.execute("delete from ref.venues")

        extras.execute_values(cur, """
            insert into ref.venues (venue_code, name, venue_type, address)
            values %s""", venues)

        extras.execute_values(cur, """
            insert into ref.suppliers (supplier_code, name, category, billing_email,
                payment_terms, layout, billing_mode, invoice_no_pattern, status)
            values %s""", suppliers)

        extras.execute_values(cur, """
            insert into ref.catalog_items (sku, supplier_code, description, unit,
                alt_unit, alt_unit_factor, vat_rate, discount_threshold_qty, discount_pct)
            values %s""", catalog)

        extras.execute_values(cur, """
            insert into ref.price_versions (version_code, effective_from, effective_to,
                scope_supplier, note)
            values %s""", versions)

        # Здесь сработает ограничение непересечения, если в прайсе две цены
        # на одну дату. Это то место, где такая ошибка обязана всплыть.
        extras.execute_values(cur, """
            insert into ref.prices (sku, version_code, unit, unit_price, valid_from, valid_to)
            values %s""", prices)

        extras.execute_values(cur, """
            insert into truth.documents (doc_id, invoice_number, supplier, venue, layout,
                received_date, line_count, total_eur, file_kind, readability, softened,
                defects, duplicate_of, duplicate_kind)
            values %s""", documents, page_size=100)

        extras.execute_values(cur, """
            insert into truth.defects (doc_id, invoice_number, defect_type, supplier, venue,
                layout, scan_kind, invoice_date, line_no, field, correct_value,
                invoiced_value, impact_eur, note)
            values %s""", defects, page_size=100)

        extras.execute_values(cur, """
            insert into truth.batch (key, value) values %s""", batch)


def report(price_data, batch_data):
    venues, suppliers, catalog, versions, prices = price_data
    documents, defects, batch, _ = batch_data

    expected = {
        "ref.venues": (len(venues), 4),
        "ref.suppliers": (len(suppliers), 12),
        "ref.catalog_items": (len(catalog), 104),
        "ref.price_versions": (len(versions), 3),
        "ref.prices": (len(prices), 216),
        "truth.documents": (len(documents), 260),
        "truth.defects": (len(defects), 98),
    }

    print()
    ok = True
    for name, (got, want) in expected.items():
        mark = "OK " if got == want else "!! "
        if got != want:
            ok = False
        print(f"  {mark}{name:<20} {got:>4}  ожидалось {want}")

    print()
    for code, _, _, scope, _ in versions:
        print(f"  версия {code:<16} область действия: {scope or 'все поставщики'}")

    types = {}
    for d in defects:
        types[d[2]] = types.get(d[2], 0) + 1
    print()
    for t, n in sorted(types.items(), key=lambda x: (-x[1], x[0])):
        print(f"  {t:<20} {n:>3}")

    dup = sum(1 for d in documents if d[13])
    print(f"\n  дублей в партии: {dup}  (ожидалось 12)")
    if dup != 12:
        ok = False

    return ok


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--price", required=True, help="kestrel-price-list.xlsx")
    ap.add_argument("--batch", required=True, help="kestrel-april-2026-batch.xlsx")
    ap.add_argument("--dry-run", action="store_true", help="только разбор, без базы")
    args = ap.parse_args()

    for p in (args.price, args.batch):
        if not os.path.exists(p):
            raise SystemExit(f"Файл не найден: {p}")

    print(f"Читаю {args.batch}")
    batch_data = read_batch(args.batch)
    print(f"Читаю {args.price}")
    price_data = read_price_list(args.price, batch_data[3])

    ok = report(price_data, batch_data)
    if not ok:
        raise SystemExit("\nКоличества не сошлись. Скорее всего это другая версия файла — "
                         "сверьтесь с признаками актуальных версий в 00-opis.md.")

    if args.dry_run:
        print("\nРазбор прошёл. База не тронута.")
        return

    url = os.environ.get("DATABASE_URL")
    if not url:
        raise SystemExit("Не задана переменная DATABASE_URL")

    import psycopg2

    print("\nПодключаюсь")
    conn = psycopg2.connect(url)
    try:
        load(conn, price_data, batch_data)
        conn.commit()
        print("Загружено. Теперь запускайте 99-proverka.sql")
    except Exception:
        conn.rollback()
        print("Откат: база не изменилась", file=sys.stderr)
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    main()
